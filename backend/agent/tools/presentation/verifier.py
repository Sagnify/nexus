"""Verification logic for generated PowerPoint (.pptx) presentations."""
from __future__ import annotations

import logging
from pathlib import Path
from typing import Tuple, Dict, Any, Optional
import pptx
from pptx import Presentation

from backend.core.paths import resolve_system_path

logger = logging.getLogger("nexus.presentation_verifier")


def verify_presentation(
    file_path: str | Path,
    min_slides: int = 1,
    required_topic: Optional[str] = None,
) -> Tuple[bool, Dict[str, Any], str]:
    """
    Verify that a generated PowerPoint presentation exists, is non-corrupt,
    has sufficient slide count and text content.
    """
    try:
        resolved = resolve_system_path(file_path)
        if not resolved.exists():
            return False, {}, f"Presentation file does not exist at '{resolved}'"

        if resolved.stat().st_size == 0:
            return False, {}, f"Presentation file at '{resolved}' is empty (0 bytes)."

        prs = Presentation(str(resolved))
        slide_count = len(prs.slides)

        if slide_count < min_slides:
            return False, {"slide_count": slide_count}, f"Presentation has {slide_count} slides, expected at least {min_slides}."

        slide_titles = []
        total_text_length = 0

        for slide in prs.slides:
            first_text = ""
            for shape in slide.shapes:
                if shape.has_text_frame and shape.text_frame.text:
                    txt = shape.text_frame.text.strip()
                    total_text_length += len(txt)
                    if not first_text and txt:
                        first_text = txt.split("\n")[0].strip()
            slide_titles.append(first_text or "Untitled Slide")

        if total_text_length < 50:
            return False, {"slide_count": slide_count, "text_length": total_text_length}, "Presentation text content appears too sparse or empty."

        stats = {
            "slide_count": slide_count,
            "slide_titles": slide_titles,
            "file_size_bytes": resolved.stat().st_size,
            "resolved_path": str(resolved),
        }

        return True, stats, ""

    except Exception as err:
        logger.error(f"Presentation verification failed: {err}")
        return False, {}, f"Failed to verify presentation: {err}"
