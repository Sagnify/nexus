"""PowerPoint presentation adapter using python-pptx with professional 16:9 themes."""
from __future__ import annotations

import asyncio
import datetime
import json
import logging
import os
import re
import subprocess
import sys
import uuid
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple
from PIL import Image

import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

from backend.agent.tools.presentation.models import (
    SlideItem,
    PresentationData,
    PresentationStats,
)
from backend.core.paths import (
    get_active_desktop,
    get_active_documents,
    get_active_downloads,
    resolve_system_path,
)

logger = logging.getLogger("nexus.presentation_adapter")

# Theme Color Palettes (RGB)
THEMES = {
    "executive_navy": {
        "name": "Executive Navy",
        "bg": RGBColor(15, 23, 42),          # #0F172A
        "card_bg": RGBColor(30, 41, 59),     # #1E293B
        "accent": RGBColor(56, 189, 248),     # Sky 400 (#38BDF8)
        "accent_secondary": RGBColor(129, 140, 248),  # Indigo 400 (#818CF8)
        "text_primary": RGBColor(248, 250, 252),  # #F8FAFC
        "text_secondary": RGBColor(148, 163, 184), # #94A3B8
        "border": RGBColor(51, 65, 85),      # #334155
        "card_highlight": RGBColor(14, 165, 233),
    },
    "modern_dark": {
        "name": "Modern Dark",
        "bg": RGBColor(18, 18, 20),          # #121214
        "card_bg": RGBColor(30, 30, 36),     # #1E1E24
        "accent": RGBColor(168, 85, 247),    # Purple 500 (#A855F7)
        "accent_secondary": RGBColor(6, 182, 212), # Cyan 500 (#06B6D4)
        "text_primary": RGBColor(255, 255, 255),
        "text_secondary": RGBColor(161, 161, 170), # #A1A1AA
        "border": RGBColor(39, 39, 42),      # #27272A
        "card_highlight": RGBColor(147, 51, 234),
    },
    "tech_indigo": {
        "name": "Tech Indigo",
        "bg": RGBColor(10, 15, 29),          # #0A0F1D
        "card_bg": RGBColor(17, 28, 51),     # #111C33
        "accent": RGBColor(99, 102, 241),    # Indigo 500 (#6366F1)
        "accent_secondary": RGBColor(16, 185, 129), # Emerald 500 (#10B981)
        "text_primary": RGBColor(241, 245, 249),
        "text_secondary": RGBColor(148, 163, 184),
        "border": RGBColor(30, 41, 59),
        "card_highlight": RGBColor(79, 70, 229),
    },
    "emerald_green": {
        "name": "Emerald Crisp",
        "bg": RGBColor(6, 78, 59),           # Deep Forest #064E3B
        "card_bg": RGBColor(4, 120, 87),     # #047857
        "accent": RGBColor(52, 211, 153),    # Emerald 400 #34D399
        "accent_secondary": RGBColor(251, 191, 36), # Amber 400 #FBBF24
        "text_primary": RGBColor(255, 255, 255),
        "text_secondary": RGBColor(209, 250, 229),
        "border": RGBColor(16, 185, 129),
        "card_highlight": RGBColor(5, 150, 105),
    },
    "corporate_light": {
        "name": "Corporate Light",
        "bg": RGBColor(248, 250, 252),       # #F8FAFC
        "card_bg": RGBColor(255, 255, 255),  # Pure White
        "accent": RGBColor(37, 99, 235),     # Blue 600 (#2563EB)
        "accent_secondary": RGBColor(13, 148, 136), # Teal 600 (#0D9488)
        "text_primary": RGBColor(15, 23, 42), # Dark Slate #0F172A
        "text_secondary": RGBColor(71, 85, 105), # #475569
        "border": RGBColor(226, 232, 240),   # #E2E8F0
        "card_highlight": RGBColor(29, 78, 216),
    },
}

FONT_TITLE = "Segoe UI"
FONT_BODY = "Segoe UI"


class PresentationAdapter:
    """Orchestrates PowerPoint (.pptx) creation, slide generation, and storage."""

    def __init__(self):
        self._sessions: Dict[str, Dict[str, Any]] = {}

    def _get_theme(self, theme_name: Optional[str]) -> Dict[str, Any]:
        key = str(theme_name or "executive_navy").lower().replace(" ", "_")
        return THEMES.get(key, THEMES["executive_navy"])

    def _calculate_image_fit(self, image_path: Path | str, max_w_in: float, max_h_in: float) -> Tuple[float, float]:
        """Calculate width and height in inches preserving aspect ratio within bounding box."""
        try:
            with Image.open(str(image_path)) as img:
                orig_w, orig_h = img.size
            if orig_w <= 0 or orig_h <= 0:
                return max_w_in, max_h_in
            aspect = orig_w / orig_h
            box_aspect = max_w_in / max_h_in
            if aspect > box_aspect:
                w = max_w_in
                h = max_w_in / aspect
            else:
                h = max_h_in
                w = max_h_in * aspect
            return round(w, 3), round(h, 3)
        except Exception as e:
            logger.debug(f"Failed to calculate image dimensions for {image_path}: {e}")
            return max_w_in, max_h_in

    def _apply_slide_background(self, slide, theme: Dict[str, Any], width_in: float = 13.333, height_in: float = 7.5):
        """Paint full-bleed background color shape."""
        shapes = slide.shapes
        bg_shape = shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(0), Inches(0), Inches(width_in), Inches(height_in)
        )
        bg_shape.fill.solid()
        bg_shape.fill.fore_color.rgb = theme["bg"]
        bg_shape.line.fill.background()  # No border

    def _render_header(self, slide, title_text: str, theme: Dict[str, Any], category: str = "PRESENTATION"):
        """Render a consistent, clean slide header with category tag and accent underline."""
        shapes = slide.shapes

        # Category eyebrow
        cat_box = shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.7), Inches(0.4))
        tf_cat = cat_box.text_frame
        tf_cat.word_wrap = True
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = category.upper()
        p_cat.font.name = FONT_BODY
        p_cat.font.size = Pt(10)
        p_cat.font.bold = True
        p_cat.font.color.rgb = theme["accent"]

        # Main slide title
        title_box = shapes.add_textbox(Inches(0.8), Inches(0.8), Inches(11.7), Inches(0.8))
        tf = title_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title_text
        p.font.name = FONT_TITLE
        p.font.size = Pt(26)
        p.font.bold = True
        p.font.color.rgb = theme["text_primary"]

        # Sleek accent line
        line = shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(0.8), Inches(1.65), Inches(1.8), Inches(0.04)
        )
        line.fill.solid()
        line.fill.fore_color.rgb = theme["accent"]
        line.line.fill.background()

    def _render_footer(self, slide, theme: Dict[str, Any], slide_num: int, total_slides: int, footer_text: str = "NEXUS AI"):
        """Render footer bar with slide numbers and branding."""
        shapes = slide.shapes
        box = shapes.add_textbox(Inches(0.8), Inches(6.9), Inches(11.7), Inches(0.4))
        tf = box.text_frame
        p = tf.paragraphs[0]
        p.text = f"{footer_text}   •   Slide {slide_num} of {total_slides}"
        p.font.name = FONT_BODY
        p.font.size = Pt(9)
        p.font.color.rgb = theme["text_secondary"]

    def _render_title_slide(self, prs, slide_data: SlideItem, theme: Dict[str, Any], total_slides: int):
        blank_layout = prs.slide_layouts[6]
        slide = prs.slides.add_slide(blank_layout)
        self._apply_slide_background(slide, theme)

        shapes = slide.shapes
        has_image = bool(slide_data.image_path and Path(slide_data.image_path).exists())

        if has_image:
            # Executive Split Title Layout (Left: Text & Metadata, Right: Hero Image Card)
            bar = shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(0.12), Inches(3.6))
            bar.fill.solid()
            bar.fill.fore_color.rgb = theme["accent"]
            bar.line.fill.background()

            tag_box = shapes.add_textbox(Inches(1.1), Inches(1.6), Inches(5.8), Inches(0.35))
            p_tag = tag_box.text_frame.paragraphs[0]
            p_tag.text = "EXECUTIVE BRIEFING & PRESENTATION"
            p_tag.font.name = FONT_BODY
            p_tag.font.size = Pt(10)
            p_tag.font.bold = True
            p_tag.font.color.rgb = theme["accent"]

            title_box = shapes.add_textbox(Inches(1.1), Inches(2.0), Inches(5.8), Inches(2.0))
            tf = title_box.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = slide_data.title
            p.font.name = FONT_TITLE
            p.font.size = Pt(32)
            p.font.bold = True
            p.font.color.rgb = theme["text_primary"]

            sub_text = slide_data.subtitle or "Comprehensive Analysis & Strategic Overview"
            sub_box = shapes.add_textbox(Inches(1.1), Inches(4.1), Inches(5.8), Inches(0.9))
            tf_sub = sub_box.text_frame
            tf_sub.word_wrap = True
            p_sub = tf_sub.paragraphs[0]
            p_sub.text = sub_text
            p_sub.font.name = FONT_BODY
            p_sub.font.size = Pt(14)
            p_sub.font.color.rgb = theme["text_secondary"]

            pill = shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.1), Inches(5.2), Inches(4.8), Inches(0.5))
            pill.fill.solid()
            pill.fill.fore_color.rgb = theme["card_bg"]
            pill.line.color.rgb = theme["border"]
            p_pill = pill.text_frame.paragraphs[0]
            p_pill.text = f"Prepared by NEXUS AI   •   {datetime.datetime.now().strftime('%B %Y')}"
            p_pill.font.name = FONT_BODY
            p_pill.font.size = Pt(10)
            p_pill.font.color.rgb = theme["accent_secondary"]
            p_pill.alignment = PP_ALIGN.LEFT

            # Right Hero Image Card Container
            container_x = 7.2
            container_y = 1.6
            container_w = 5.3
            container_h = 4.8
            img_card = shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE,
                Inches(container_x), Inches(container_y), Inches(container_w), Inches(container_h)
            )
            img_card.fill.solid()
            img_card.fill.fore_color.rgb = theme["card_bg"]
            img_card.line.color.rgb = theme["border"]

            fit_w, fit_h = self._calculate_image_fit(slide_data.image_path, max_w_in=5.0, max_h_in=4.1)
            pic_x = container_x + (container_w - fit_w) / 2
            pic_y = container_y + 0.15 + (4.1 - fit_h) / 2
            slide.shapes.add_picture(str(slide_data.image_path), Inches(pic_x), Inches(pic_y), Inches(fit_w), Inches(fit_h))

            if slide_data.image_caption:
                cap_box = shapes.add_textbox(Inches(container_x + 0.2), Inches(container_y + 4.35), Inches(container_w - 0.4), Inches(0.35))
                p_cap = cap_box.text_frame.paragraphs[0]
                p_cap.text = f"Photo: {slide_data.image_caption}"
                p_cap.font.name = FONT_BODY
                p_cap.font.size = Pt(9)
                p_cap.font.color.rgb = theme["text_secondary"]
                p_cap.alignment = PP_ALIGN.CENTER
        else:
            # Full-Width Title Layout
            bar = shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(2.2), Inches(0.12), Inches(2.8))
            bar.fill.solid()
            bar.fill.fore_color.rgb = theme["accent"]
            bar.line.fill.background()

            tag_box = shapes.add_textbox(Inches(1.3), Inches(2.0), Inches(10.5), Inches(0.4))
            p_tag = tag_box.text_frame.paragraphs[0]
            p_tag.text = "EXECUTIVE BRIEFING & PRESENTATION"
            p_tag.font.name = FONT_BODY
            p_tag.font.size = Pt(11)
            p_tag.font.bold = True
            p_tag.font.color.rgb = theme["accent"]

            title_box = shapes.add_textbox(Inches(1.3), Inches(2.4), Inches(10.5), Inches(1.6))
            tf = title_box.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = slide_data.title
            p.font.name = FONT_TITLE
            p.font.size = Pt(40)
            p.font.bold = True
            p.font.color.rgb = theme["text_primary"]

            sub_text = slide_data.subtitle or "Comprehensive Analysis & Strategic Overview"
            sub_box = shapes.add_textbox(Inches(1.3), Inches(4.1), Inches(10.5), Inches(0.8))
            tf_sub = sub_box.text_frame
            tf_sub.word_wrap = True
            p_sub = tf_sub.paragraphs[0]
            p_sub.text = sub_text
            p_sub.font.name = FONT_BODY
            p_sub.font.size = Pt(18)
            p_sub.font.color.rgb = theme["text_secondary"]

            pill = shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.3), Inches(5.1), Inches(4.5), Inches(0.55))
            pill.fill.solid()
            pill.fill.fore_color.rgb = theme["card_bg"]
            pill.line.color.rgb = theme["border"]
            p_pill = pill.text_frame.paragraphs[0]
            p_pill.text = f"Prepared by NEXUS AI   •   {datetime.datetime.now().strftime('%B %Y')}"
            p_pill.font.name = FONT_BODY
            p_pill.font.size = Pt(10)
            p_pill.font.color.rgb = theme["accent_secondary"]
            p_pill.alignment = PP_ALIGN.LEFT

        self._render_footer(slide, theme, 1, total_slides)

    def _render_agenda_slide(self, prs, slide_data: SlideItem, theme: Dict[str, Any], slide_num: int, total_slides: int):
        blank_layout = prs.slide_layouts[6]
        slide = prs.slides.add_slide(blank_layout)
        self._apply_slide_background(slide, theme)
        self._render_header(slide, slide_data.title or "Executive Agenda & Roadmap", theme, category="OVERVIEW")

        items = slide_data.bullets if slide_data.bullets else [
            "Executive Summary & Strategic Context",
            "Key Industry Drivers & Technical Foundations",
            "Core Architecture & Practical Applications",
            "Market Opportunities & Competitive Dynamics",
            "Strategic Recommendations & Next Steps",
        ]

        # Render 2-column or list cards
        shapes = slide.shapes
        card_w = Inches(5.6)
        card_h = Inches(1.0)
        start_x_left = Inches(0.8)
        start_x_right = Inches(6.9)
        start_y = Inches(2.0)

        for idx, item in enumerate(items[:6]):
            col = idx % 2
            row = idx // 2
            x = start_x_left if col == 0 else start_x_right
            y = start_y + (row * Inches(1.4))

            card = shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, card_w, card_h)
            card.fill.solid()
            card.fill.fore_color.rgb = theme["card_bg"]
            card.line.color.rgb = theme["border"]

            tf = card.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = f"{idx + 1:02d}.   {item}"
            p.font.name = FONT_BODY
            p.font.size = Pt(14)
            p.font.bold = True
            p.font.color.rgb = theme["text_primary"]
            p.alignment = PP_ALIGN.LEFT

        self._render_footer(slide, theme, slide_num, total_slides)

    def _render_content_slide(self, prs, slide_data: SlideItem, theme: Dict[str, Any], slide_num: int, total_slides: int):
        blank_layout = prs.slide_layouts[6]
        slide = prs.slides.add_slide(blank_layout)
        self._apply_slide_background(slide, theme)
        self._render_header(slide, slide_data.title, theme, category="INSIGHTS")

        shapes = slide.shapes
        has_image = bool(slide_data.image_path and Path(slide_data.image_path).exists())

        if has_image:
            # Split Content Layout: Left Bullets (w=6.4"), Right Image (w=5.0")
            left_w = Inches(6.4)
            left_h = Inches(3.7) if slide_data.key_takeaway else Inches(4.7)

            card = shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.9), left_w, left_h)
            card.fill.solid()
            card.fill.fore_color.rgb = theme["card_bg"]
            card.line.color.rgb = theme["border"]

            tf = card.text_frame
            tf.word_wrap = True
            tf.clear()

            bullets = slide_data.bullets or ["Key development point", "Significant factor influencing outcomes"]
            for i, bullet in enumerate(bullets):
                p = tf.add_paragraph() if i > 0 else tf.paragraphs[0]
                p.text = f"•  {bullet}"
                p.font.name = FONT_BODY
                p.font.size = Pt(13)
                p.font.color.rgb = theme["text_primary"]
                p.space_before = Pt(6)
                p.space_after = Pt(6)

            if slide_data.key_takeaway:
                takeaway_box = shapes.add_shape(
                    MSO_SHAPE.ROUNDED_RECTANGLE,
                    Inches(0.8), Inches(5.8), left_w, Inches(0.9)
                )
                takeaway_box.fill.solid()
                takeaway_box.fill.fore_color.rgb = theme["bg"]
                takeaway_box.line.color.rgb = theme["accent"]

                tf_tk = takeaway_box.text_frame
                tf_tk.word_wrap = True
                p_tk = tf_tk.paragraphs[0]
                p_tk.text = f"KEY TAKEAWAY:  {slide_data.key_takeaway}"
                p_tk.font.name = FONT_BODY
                p_tk.font.size = Pt(11)
                p_tk.font.bold = True
                p_tk.font.color.rgb = theme["accent"]

            # Right Column Image Card
            right_x = 7.5
            right_y = 1.9
            right_w = 5.0
            right_h = 4.7
            img_card = shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE,
                Inches(right_x), Inches(right_y), Inches(right_w), Inches(right_h)
            )
            img_card.fill.solid()
            img_card.fill.fore_color.rgb = theme["card_bg"]
            img_card.line.color.rgb = theme["border"]

            fit_w, fit_h = self._calculate_image_fit(slide_data.image_path, max_w_in=4.7, max_h_in=4.0)
            pic_x = right_x + (right_w - fit_w) / 2
            pic_y = right_y + 0.15 + (4.0 - fit_h) / 2
            slide.shapes.add_picture(str(slide_data.image_path), Inches(pic_x), Inches(pic_y), Inches(fit_w), Inches(fit_h))

            if slide_data.image_caption:
                cap_box = shapes.add_textbox(Inches(right_x + 0.2), Inches(right_y + 4.25), Inches(right_w - 0.4), Inches(0.35))
                p_cap = cap_box.text_frame.paragraphs[0]
                p_cap.text = f"Photo: {slide_data.image_caption}"
                p_cap.font.name = FONT_BODY
                p_cap.font.size = Pt(9)
                p_cap.font.color.rgb = theme["text_secondary"]
                p_cap.alignment = PP_ALIGN.CENTER
        else:
            # Full-Width Content Card
            main_h = Inches(3.7) if slide_data.key_takeaway else Inches(4.7)
            card = shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.9), Inches(11.7), main_h)
            card.fill.solid()
            card.fill.fore_color.rgb = theme["card_bg"]
            card.line.color.rgb = theme["border"]

            tf = card.text_frame
            tf.word_wrap = True
            tf.clear()

            bullets = slide_data.bullets or ["Key development point", "Significant factor influencing outcomes"]
            for i, bullet in enumerate(bullets):
                p = tf.add_paragraph() if i > 0 else tf.paragraphs[0]
                p.text = f"•  {bullet}"
                p.font.name = FONT_BODY
                p.font.size = Pt(14)
                p.font.color.rgb = theme["text_primary"]
                p.space_before = Pt(8)
                p.space_after = Pt(8)

            if slide_data.key_takeaway:
                takeaway_box = shapes.add_shape(
                    MSO_SHAPE.ROUNDED_RECTANGLE,
                    Inches(0.8), Inches(5.8), Inches(11.7), Inches(0.9)
                )
                takeaway_box.fill.solid()
                takeaway_box.fill.fore_color.rgb = theme["bg"]
                takeaway_box.line.color.rgb = theme["accent"]

                tf_tk = takeaway_box.text_frame
                tf_tk.word_wrap = True
                p_tk = tf_tk.paragraphs[0]
                p_tk.text = f"KEY TAKEAWAY:  {slide_data.key_takeaway}"
                p_tk.font.name = FONT_BODY
                p_tk.font.size = Pt(12)
                p_tk.font.bold = True
                p_tk.font.color.rgb = theme["accent"]

        self._render_footer(slide, theme, slide_num, total_slides)

    def _render_two_column_slide(self, prs, slide_data: SlideItem, theme: Dict[str, Any], slide_num: int, total_slides: int):
        blank_layout = prs.slide_layouts[6]
        slide = prs.slides.add_slide(blank_layout)
        self._apply_slide_background(slide, theme)
        self._render_header(slide, slide_data.title, theme, category="COMPARATIVE ANALYSIS")

        shapes = slide.shapes
        col_w = Inches(5.7)
        col_h = Inches(4.6)

        # Left Column Card
        left_card = shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(2.0), col_w, col_h)
        left_card.fill.solid()
        left_card.fill.fore_color.rgb = theme["card_bg"]
        left_card.line.color.rgb = theme["border"]

        tf_l = left_card.text_frame
        tf_l.word_wrap = True
        p_l_head = tf_l.paragraphs[0]
        p_l_head.text = (slide_data.left_heading or "Key Advantages & Strengths").upper()
        p_l_head.font.name = FONT_TITLE
        p_l_head.font.size = Pt(14)
        p_l_head.font.bold = True
        p_l_head.font.color.rgb = theme["accent"]
        p_l_head.space_after = Pt(12)

        left_bullets = slide_data.left_bullets or slide_data.bullets[:3]
        for b in left_bullets:
            p = tf_l.add_paragraph()
            p.text = f"✓  {b}"
            p.font.name = FONT_BODY
            p.font.size = Pt(13)
            p.font.color.rgb = theme["text_primary"]
            p.space_after = Pt(8)

        # Right Column Card
        right_card = shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(2.0), col_w, col_h)
        right_card.fill.solid()
        right_card.fill.fore_color.rgb = theme["card_bg"]
        right_card.line.color.rgb = theme["border"]

        tf_r = right_card.text_frame
        tf_r.word_wrap = True
        p_r_head = tf_r.paragraphs[0]
        p_r_head.text = (slide_data.right_heading or "Challenges & Critical Considerations").upper()
        p_r_head.font.name = FONT_TITLE
        p_r_head.font.size = Pt(14)
        p_r_head.font.bold = True
        p_r_head.font.color.rgb = theme["accent_secondary"]
        p_r_head.space_after = Pt(12)

        right_bullets = slide_data.right_bullets or slide_data.bullets[3:]
        for b in right_bullets:
            p = tf_r.add_paragraph()
            p.text = f"•  {b}"
            p.font.name = FONT_BODY
            p.font.size = Pt(13)
            p.font.color.rgb = theme["text_primary"]
            p.space_after = Pt(8)

        self._render_footer(slide, theme, slide_num, total_slides)

    def _render_stats_slide(self, prs, slide_data: SlideItem, theme: Dict[str, Any], slide_num: int, total_slides: int):
        blank_layout = prs.slide_layouts[6]
        slide = prs.slides.add_slide(blank_layout)
        self._apply_slide_background(slide, theme)
        self._render_header(slide, slide_data.title, theme, category="KEY METRICS")

        shapes = slide.shapes
        stats = slide_data.stats or [
            {"value": "85%", "label": "Efficiency Gain"},
            {"value": "3.5x", "label": "Adoption Rate"},
            {"value": "$1.2T", "label": "Projected Market Size"},
        ]

        count = len(stats)
        spacing = Inches(0.4)
        total_avail = Inches(11.7)
        card_w = (total_avail - (spacing * (count - 1))) / count
        card_h = Inches(3.2)
        start_x = Inches(0.8)
        y = Inches(2.2)

        for idx, item in enumerate(stats):
            x = start_x + (idx * (card_w + spacing))
            card = shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, card_w, card_h)
            card.fill.solid()
            card.fill.fore_color.rgb = theme["card_bg"]
            card.line.color.rgb = theme["border"]

            tf = card.text_frame
            tf.word_wrap = True

            p_val = tf.paragraphs[0]
            p_val.text = str(item.get("value", ""))
            p_val.font.name = FONT_TITLE
            p_val.font.size = Pt(36)
            p_val.font.bold = True
            p_val.font.color.rgb = theme["accent"]
            p_val.alignment = PP_ALIGN.CENTER
            p_val.space_before = Pt(30)
            p_val.space_after = Pt(10)

            p_lbl = tf.add_paragraph()
            p_lbl.text = str(item.get("label", ""))
            p_lbl.font.name = FONT_BODY
            p_lbl.font.size = Pt(13)
            p_lbl.font.color.rgb = theme["text_secondary"]
            p_lbl.alignment = PP_ALIGN.CENTER

        if slide_data.key_takeaway:
            takeaway_box = shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE,
                Inches(0.8), Inches(5.8), Inches(11.7), Inches(0.8)
            )
            takeaway_box.fill.solid()
            takeaway_box.fill.fore_color.rgb = theme["bg"]
            takeaway_box.line.color.rgb = theme["accent"]
            p_tk = takeaway_box.text_frame.paragraphs[0]
            p_tk.text = f"SUMMARY: {slide_data.key_takeaway}"
            p_tk.font.name = FONT_BODY
            p_tk.font.size = Pt(12)
            p_tk.font.bold = True
            p_tk.font.color.rgb = theme["accent"]

        self._render_footer(slide, theme, slide_num, total_slides)

    def _render_conclusion_slide(self, prs, slide_data: SlideItem, theme: Dict[str, Any], slide_num: int, total_slides: int):
        blank_layout = prs.slide_layouts[6]
        slide = prs.slides.add_slide(blank_layout)
        self._apply_slide_background(slide, theme)
        self._render_header(slide, slide_data.title or "Summary & Strategic Roadmap", theme, category="CONCLUSION")

        shapes = slide.shapes

        # Summary box
        card = shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(2.0), Inches(11.7), Inches(3.4))
        card.fill.solid()
        card.fill.fore_color.rgb = theme["card_bg"]
        card.line.color.rgb = theme["border"]

        tf = card.text_frame
        tf.word_wrap = True
        bullets = slide_data.bullets or [
            "Actionable insights synthesized across core operational dimensions",
            "Clear milestone prioritization establishes measurable value realization",
            "Cross-functional alignment ensures accelerated deployment with minimal friction",
        ]
        for i, b in enumerate(bullets):
            p = tf.add_paragraph() if i > 0 else tf.paragraphs[0]
            p.text = f"★  {b}"
            p.font.name = FONT_BODY
            p.font.size = Pt(14)
            p.font.color.rgb = theme["text_primary"]
            p.space_after = Pt(10)

        # Call to Action card
        cta = shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(5.6), Inches(11.7), Inches(1.0))
        cta.fill.solid()
        cta.fill.fore_color.rgb = theme["bg"]
        cta.line.color.rgb = theme["accent"]

        tf_cta = cta.text_frame
        tf_cta.word_wrap = True
        p_cta = tf_cta.paragraphs[0]
        p_cta.text = "QUESTIONS & DISCUSSION"
        p_cta.font.name = FONT_TITLE
        p_cta.font.size = Pt(14)
        p_cta.font.bold = True
        p_cta.font.color.rgb = theme["accent"]

        p_cta_sub = tf_cta.add_paragraph()
        p_cta_sub.text = slide_data.key_takeaway or "Thank you. Open for technical inquiry and strategic Q&A."
        p_cta_sub.font.name = FONT_BODY
        p_cta_sub.font.size = Pt(12)
        p_cta_sub.font.color.rgb = theme["text_secondary"]

        self._render_footer(slide, theme, slide_num, total_slides)

    async def generate_presentation_data_with_llm(self, topic: str, num_slides: int = 6, theme: str = "executive_navy") -> Optional[PresentationData]:
        """Use LLM to generate rich, in-depth presentation structure."""
        from langchain_core.messages import SystemMessage, HumanMessage
        from backend.agent.router.model_router import ainvoke_with_dynamic_switch
        from backend.core.config import get_groq_api_key, get_gemma_api_key

        if not (get_groq_api_key() or get_gemma_api_key()):
            return None

        system_prompt = (
            "You are an Elite Executive Presentation Designer and Strategic Storyteller. "
            "Your task is to generate a comprehensive, visually compelling, authoritative PowerPoint slide deck "
            "on the user's requested topic.\n\n"
            "Rules:\n"
            "1. Output STRICTLY valid JSON conforming to the PresentationData schema.\n"
            "2. Make bullet points informative, concrete, and deeply relevant with real terminology, metrics, and actionable concepts.\n"
            "3. Structure the deck with:\n"
            "   - Slide 1: 'title' slide\n"
            "   - Slide 2: 'agenda' slide\n"
            "   - Slide 3: 'content' or 'two_column' slide\n"
            "   - Slide 4: 'two_column' (strengths/challenges) or 'stats' slide\n"
            "   - Slide 5: 'content' or 'stats' slide (trends/future)\n"
            "   - Slide 6: 'conclusion' slide\n"
            "4. Never output placeholders like 'Bullet 1' or 'Point 2'. Write thorough, high-impact content.\n\n"
            "JSON Format:\n"
            "{\n"
            '  "title": "Main Presentation Title",\n'
            '  "subtitle": "Informative Subtitle",\n'
            '  "topic": "Clean topic name",\n'
            '  "theme": "executive_navy",\n'
            '  "slides": [\n'
            '    {\n'
            '      "title": "Title of Slide",\n'
            '      "slide_type": "title|agenda|content|two_column|stats|conclusion",\n'
            '      "bullets": ["Point 1...", "Point 2...", "Point 3..."],\n'
            '      "left_heading": "Optional heading for left column",\n'
            '      "left_bullets": ["Left point 1...", "Left point 2..."],\n'
            '      "right_heading": "Optional heading for right column",\n'
            '      "right_bullets": ["Right point 1...", "Right point 2..."],\n'
            '      "stats": [{"value": "85%", "label": "Metric Description"}],\n'
            '      "key_takeaway": "Key punchy insight for this slide"\n'
            '    }\n'
            '  ]\n'
            "}"
        )

        user_prompt = f"Create a {num_slides}-slide executive presentation on: {topic}. Theme: {theme}"

        try:
            res = await asyncio.wait_for(
                ainvoke_with_dynamic_switch(
                    [SystemMessage(content=system_prompt), HumanMessage(content=user_prompt)],
                    operation="reasoning",
                    temperature=0.2,
                ),
                timeout=40.0,
            )
            raw = res.content.strip()
            if "```" in raw:
                raw = re.sub(r"```(?:json)?", "", raw).replace("```", "").strip()
            if "{" in raw and "}" in raw:
                start = raw.index("{")
                end = raw.rindex("}") + 1
                slice_raw = raw[start:end]
                try:
                    data_dict = json.loads(slice_raw)
                except Exception:
                    cleaned = re.sub(r",\s*([\]}])", r"\1", slice_raw)
                    data_dict = json.loads(cleaned)
                return PresentationData(**data_dict)
            return None
        except Exception as err:
            logger.warning(f"LLM presentation data generation fallback to synthesis: {err}")
            return None


    def synthesize_fallback_presentation_data(self, topic: str, theme: str = "executive_navy") -> PresentationData:
        """High-quality deterministic presentation generator for any topic."""
        clean_topic = topic.strip().title()
        words = clean_topic.split()
        short_title = clean_topic if len(words) <= 5 else " ".join(words[:5])

        slides = [
            SlideItem(
                title=clean_topic,
                subtitle="Strategic Overview, Industry Drivers & Practical Implementation",
                slide_type="title",
                key_takeaway="Executive briefing prepared by NEXUS AI",
                image_query=clean_topic,
            ),
            SlideItem(
                title="Executive Agenda & Scope",
                slide_type="agenda",
                bullets=[
                    f"Foundations & Background of {short_title}",
                    "Key Technological & Market Drivers",
                    "Core Architectural & Operational Paradigms",
                    "Comparative Advantages & Implementation Challenges",
                    "Measurable Impact, Metrics & Benchmarks",
                    "Future Outlook & Strategic Roadmap",
                ],
            ),
            SlideItem(
                title=f"Core Foundations of {short_title}",
                slide_type="content",
                bullets=[
                    f"Rapid acceleration across global ecosystems is redefining standard approaches to {short_title.lower()}.",
                    "Integration of modernized methodologies delivers substantial efficiency and high-fidelity output.",
                    "Scalable infrastructure enables continuous adaptation while minimizing overhead and technical debt.",
                    "Cross-disciplinary convergence bridges legacy workflows with state-of-the-art automated systems.",
                ],
                key_takeaway=f"{short_title} represents a fundamental transition toward autonomous, resilient operational models.",
                image_query=f"{short_title} technology",
            ),
            SlideItem(
                title="Strategic Analysis: Opportunities vs. Challenges",
                slide_type="two_column",
                left_heading="Strategic Advantages & Drivers",
                left_bullets=[
                    "Accelerated speed-to-market and streamlined execution cycles",
                    "Significantly reduced margin of human error across complex pipelines",
                    "Data-driven visibility providing predictive rather than reactive insights",
                ],
                right_heading="Critical Risks & Constraints",
                right_bullets=[
                    "Governance, compliance, and rigorous data protection standards",
                    "Integration friction with heterogeneous legacy environments",
                    "Need for specialized expertise and organizational change management",
                ],
            ),
            SlideItem(
                title="Key Metrics & Measurable Impact",
                slide_type="stats",
                stats=[
                    {"value": "10x", "label": "Deployment Velocity"},
                    {"value": "74%", "label": "Operational Efficiency"},
                    {"value": "99.8%", "label": "Reliability & Uptime"},
                ],
                key_takeaway="Quantitative verification confirms substantial ROI and accelerated cycle completion across all tiers.",
            ),
            SlideItem(
                title="Strategic Roadmap & Recommendations",
                slide_type="conclusion",
                bullets=[
                    "Phase 1: Establish foundational baseline benchmarks and align stakeholder governance.",
                    "Phase 2: Pilot high-impact modular deployments with continuous telemetry tracking.",
                    "Phase 3: Scale enterprise-wide adoption with automated auditing and optimization loops.",
                ],
                key_takeaway=f"Accelerating {short_title} initiatives unlocks durable competitive advantage and scalable growth.",
            ),
        ]

        return PresentationData(
            title=clean_topic,
            subtitle="Strategic Overview & Executive Insights",
            topic=clean_topic,
            theme=theme,
            slides=slides,
        )

    def render_presentation(self, data: PresentationData) -> Presentation:
        """Render a PresentationData object into a python-pptx Presentation instance (16:9)."""
        prs = Presentation()
        # Set 16:9 widescreen dimensions
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.5)

        theme = self._get_theme(data.theme)
        total_slides = len(data.slides)

        for idx, slide_item in enumerate(data.slides):
            slide_num = idx + 1
            stype = slide_item.slide_type

            if stype == "title" or idx == 0:
                self._render_title_slide(prs, slide_item, theme, total_slides)
            elif stype == "agenda":
                self._render_agenda_slide(prs, slide_item, theme, slide_num, total_slides)
            elif stype == "two_column" or stype == "comparison":
                self._render_two_column_slide(prs, slide_item, theme, slide_num, total_slides)
            elif stype == "stats":
                self._render_stats_slide(prs, slide_item, theme, slide_num, total_slides)
            elif stype == "conclusion":
                self._render_conclusion_slide(prs, slide_item, theme, slide_num, total_slides)
            else:
                self._render_content_slide(prs, slide_item, theme, slide_num, total_slides)

        return prs

    async def _enrich_with_web_images(self, data: PresentationData):
        """Search and download relevant pictures from the web to embed in appropriate slides."""
        from backend.agent.tools.presentation.image_search import fetch_relevant_image

        clean_topic = re.sub(
            r"\b(presentation|powerpoint|deck|overview|slides|summary|briefing|report|intro)\b",
            "",
            data.topic,
            flags=re.IGNORECASE
        ).strip()
        if not clean_topic:
            clean_topic = data.topic

        used_paths = set()

        # 1. Title slide hero image
        if data.slides and not data.slides[0].image_path:
            hero_q = data.slides[0].image_query or clean_topic
            try:
                hero_img = await fetch_relevant_image(hero_q)
                if hero_img:
                    data.slides[0].image_path = str(hero_img)
                    data.slides[0].image_caption = clean_topic.title()
                    used_paths.add(str(hero_img))
            except Exception as e:
                logger.warning(f"Could not fetch title hero image for '{hero_q}': {e}")

        # 2. Content slides image enrichment
        for slide in data.slides[1:]:
            if slide.slide_type not in ("content", "two_column"):
                continue
            if slide.image_path:
                used_paths.add(str(slide.image_path))
                continue
            # Keep at most 2 additional content pictures so deck remains balanced and uncluttered
            if len(used_paths) >= 3:
                break

            # Derive search query for this slide
            query = slide.image_query
            if not query:
                # Remove common slide filler words to extract key noun phrase
                cleaned_title = re.sub(
                    r"\b(strategic|analysis|overview|foundations|core|impact|metrics|roadmap|recommendations|key|the|of|and|in|a|an)\b",
                    "",
                    slide.title,
                    flags=re.IGNORECASE
                ).strip()
                cleaned_title = re.sub(r"\s+", " ", cleaned_title)
                if len(cleaned_title.split()) >= 1 and len(cleaned_title) > 3:
                    query = f"{clean_topic} {cleaned_title}".strip()
                else:
                    query = clean_topic

            try:
                img_path = await fetch_relevant_image(query, fallback_topic=clean_topic)
                if img_path and str(img_path) not in used_paths:
                    slide.image_path = str(img_path)
                    slide.image_caption = query.title()
                    used_paths.add(str(img_path))
            except Exception as e:
                logger.debug(f"Could not fetch image for slide '{slide.title}': {e}")

    async def create_presentation(
        self,
        topic: str,
        theme: str = "executive_navy",
        num_slides: int = 6,
        session_id: Optional[str] = None,
        custom_data: Optional[Dict[str, Any]] = None,
        include_images: bool = True,
    ) -> Tuple[str, Path, PresentationData]:
        """
        Create a new presentation from topic or data.
        Returns: (session_id, staging_file_path, presentation_data)
        """
        sid = session_id or f"ppt-{uuid.uuid4().hex[:8]}"

        if custom_data:
            data = PresentationData(**custom_data)
        else:
            llm_data = await self.generate_presentation_data_with_llm(topic=topic, num_slides=num_slides, theme=theme)
            if llm_data:
                data = llm_data
            else:
                data = self.synthesize_fallback_presentation_data(topic=topic, theme=theme)

        # Search and insert relevant web photography if enabled
        if include_images and getattr(data, "include_images", True):
            try:
                await self._enrich_with_web_images(data)
            except Exception as img_err:
                logger.warning(f"Web image enrichment encountered an issue: {img_err}")

        prs = self.render_presentation(data)

        # Stage in temporary / scratch folder
        scratch_dir = Path("scratch")
        scratch_dir.mkdir(parents=True, exist_ok=True)
        sanitized_topic = re.sub(r"[^\w\s-]", "", data.topic).strip().replace(" ", "_")
        if not sanitized_topic:
            sanitized_topic = "Presentation"
        staging_filename = f"{sanitized_topic}_{sid}.pptx"
        staging_path = (scratch_dir / staging_filename).resolve()

        prs.save(str(staging_path))

        self._sessions[sid] = {
            "session_id": sid,
            "data": data,
            "presentation": prs,
            "staging_path": staging_path,
            "default_filename": f"{sanitized_topic}.pptx",
            "topic": data.topic,
        }

        logger.info(f"Created presentation session '{sid}' staged at '{staging_path}'")
        return sid, staging_path, data

    def save_session(
        self,
        session_id: Optional[str] = None,
        output_path: Optional[str] = None,
        default_filename: Optional[str] = None,
    ) -> Tuple[str, Path]:
        """
        Save the presentation from session to the user's chosen output destination.
        """
        # Pick session
        if session_id and session_id in self._sessions:
            sess = self._sessions[session_id]
        elif self._sessions:
            sess = list(self._sessions.values())[-1]
            session_id = sess["session_id"]
        else:
            raise ValueError("No active presentation session to save.")

        def_name = default_filename or sess.get("default_filename") or "presentation.pptx"
        if not def_name.lower().endswith(".pptx"):
            def_name += ".pptx"

        if not output_path or not str(output_path).strip():
            target_file = get_active_desktop() / def_name
        else:
            target_file = resolve_system_path(output_path, default_filename=def_name)

        if target_file.is_dir() or str(output_path).lower().rstrip("/\\") in ("desktop", "documents", "downloads"):
            target_file = target_file / def_name

        target_file.parent.mkdir(parents=True, exist_ok=True)

        prs = sess["presentation"]
        prs.save(str(target_file))

        sess["saved_path"] = target_file
        logger.info(f"Saved presentation session '{session_id}' to '{target_file}'")

        return session_id, target_file

    def reveal_in_explorer(self, file_path: Path | str) -> bool:
        """Open Windows File Explorer with the saved presentation highlighted."""
        p = Path(file_path).resolve()
        if not p.exists():
            return False

        if sys.platform == "win32":
            try:
                # Use explorer.exe /select,"<path>"
                subprocess.Popen(["explorer.exe", f"/select,{str(p)}"])
                return True
            except Exception as e:
                logger.warning(f"Could not reveal file in explorer: {e}")
        return False

    def read_session_stats(self, session_id_or_path: str) -> PresentationStats:
        """Retrieve presentation structure and stats."""
        p_path = Path(session_id_or_path)
        if p_path.exists() and p_path.is_file():
            prs = Presentation(str(p_path))
            titles = []
            for slide in prs.slides:
                title = ""
                for shape in slide.shapes:
                    if shape.has_text_frame and shape.text_frame.text:
                        title = shape.text_frame.text.split("\n")[0].strip()
                        break
                titles.append(title or "Slide")
            return PresentationStats(
                slide_count=len(prs.slides),
                title=titles[0] if titles else "Presentation",
                topic=titles[0] if titles else "Presentation",
                theme="custom",
                slide_titles=titles,
                file_path=str(p_path),
            )

        if session_id_or_path in self._sessions:
            sess = self._sessions[session_id_or_path]
            data: PresentationData = sess["data"]
            return PresentationStats(
                slide_count=len(data.slides),
                title=data.title,
                topic=data.topic,
                theme=data.theme,
                slide_titles=[s.title for s in data.slides],
                file_path=str(sess.get("saved_path") or sess.get("staging_path")),
            )

        raise ValueError(f"Presentation '{session_id_or_path}' not found.")


presentation_adapter = PresentationAdapter()
