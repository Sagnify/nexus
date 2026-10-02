"""Adapter encapsulating openpyxl operations and workbook session management."""
from __future__ import annotations
import os
import uuid
import logging
from pathlib import Path
from typing import Any, Optional

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.chart import BarChart, LineChart, PieChart, Reference

from backend.agent.tools.spreadsheet.models import SpreadsheetStats

logger = logging.getLogger("nexus.excel_adapter")


class ExcelAdapter:
    """Stateful adapter for managing XLSX workbooks and sessions via openpyxl."""

    def __init__(self):
        # Maps session_id -> {"wb": openpyxl.Workbook, "path": Path | None, "is_dirty": bool}
        self._sessions: dict[str, dict[str, Any]] = {}
        self._last_active_session_id: Optional[str] = None

    def _resolve_path(self, raw_path: str) -> Path:
        """Expand user, env vars, handle OneDrive redirection, and relative shortcut prefixes (e.g. Desktop/file.xlsx)."""
        from backend.core.paths import resolve_system_path
        return resolve_system_path(raw_path, default_filename="spreadsheet.xlsx")

    def _resolve_session_id(self, session_id: Optional[str] = None) -> str:
        """Resolve session_id, falling back to last active session or default session."""
        if session_id and session_id in self._sessions:
            return session_id
        if self._last_active_session_id and self._last_active_session_id in self._sessions:
            return self._last_active_session_id
        if self._sessions:
            return next(iter(self._sessions.keys()))
        new_id, _ = self.create_session()
        return new_id

    def _get_wb(self, session_id: Optional[str] = None) -> tuple[str, openpyxl.Workbook]:
        sid = self._resolve_session_id(session_id)
        session = self._sessions.get(sid)
        if not session or "wb" not in session:
            raise KeyError(f"Spreadsheet session '{sid}' not found or already closed.")
        return sid, session["wb"]

    def _get_ws(self, wb: openpyxl.Workbook, sheet_name: Optional[str] = None) -> openpyxl.worksheet.worksheet.Worksheet:
        if sheet_name:
            if sheet_name in wb.sheetnames:
                return wb[sheet_name]
            # Search case-insensitive
            for sname in wb.sheetnames:
                if sname.lower() == sheet_name.lower():
                    return wb[sname]
            # Create sheet if missing
            return wb.create_sheet(title=sheet_name)
        return wb.active

    def create_session(
        self,
        path: Optional[str] = None,
        initial_sheet: str = "Sheet1",
        session_id: Optional[str] = None,
    ) -> tuple[str, str]:
        """Create a new blank XLSX workbook session."""
        sid = session_id or f"sheet_{uuid.uuid4().hex[:8]}"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = initial_sheet
        target_path = self._resolve_path(path) if path else None

        self._sessions[sid] = {
            "wb": wb,
            "path": target_path,
            "is_dirty": True,
        }
        self._last_active_session_id = sid
        resolved_str = str(target_path) if target_path else ""
        logger.info(f"Created xlsx session '{sid}' (path: {resolved_str})")
        return sid, resolved_str

    def open_session(self, path: str, session_id: Optional[str] = None) -> tuple[str, str]:
        """Open an existing .xlsx file into a session."""
        target_path = self._resolve_path(path)
        if not target_path.exists():
            raise FileNotFoundError(f"Spreadsheet file not found: {target_path}")

        wb = openpyxl.load_workbook(filename=str(target_path), data_only=False)
        sid = session_id or f"sheet_{uuid.uuid4().hex[:8]}"
        self._sessions[sid] = {
            "wb": wb,
            "path": target_path,
            "is_dirty": False,
        }
        self._last_active_session_id = sid
        logger.info(f"Opened xlsx session '{sid}' from {target_path}")
        return sid, str(target_path)

    def list_sheets(self, session_id: Optional[str] = None) -> list[str]:
        """List all worksheet names in the workbook."""
        _, wb = self._get_wb(session_id)
        return list(wb.sheetnames)

    def create_sheet(self, session_id: Optional[str] = None, title: str = "Sheet2", index: Optional[int] = None) -> str:
        """Create a new worksheet in the active workbook session."""
        sid, wb = self._get_wb(session_id)
        if title in wb.sheetnames:
            title = f"{title}_{uuid.uuid4().hex[:4]}"
        if index is not None:
            wb.create_sheet(title=title, index=index)
        else:
            wb.create_sheet(title=title)
        self._sessions[sid]["is_dirty"] = True
        return sid

    def rename_sheet(self, session_id: Optional[str], old_title: str, new_title: str) -> str:
        """Rename an existing worksheet."""
        sid, wb = self._get_wb(session_id)
        if old_title in wb.sheetnames:
            wb[old_title].title = new_title
        elif wb.active and wb.active.title.lower() == old_title.lower():
            wb.active.title = new_title
        self._sessions[sid]["is_dirty"] = True
        return sid

    def read_cell(self, session_id: Optional[str], cell: str, sheet: Optional[str] = None) -> Any:
        """Read value of a single cell (e.g. 'A1')."""
        _, wb = self._get_wb(session_id)
        ws = self._get_ws(wb, sheet)
        val = ws[cell].value
        return str(val) if val is not None else ""

    def read_range(self, session_id: Optional[str], range_str: str, sheet: Optional[str] = None) -> dict[str, Any]:
        """Read values from a cell range (e.g. 'A1:C10')."""
        _, wb = self._get_wb(session_id)
        ws = self._get_ws(wb, sheet)
        cells = ws[range_str]

        rows_data = []
        if isinstance(cells, tuple):
            for row in cells:
                if isinstance(row, tuple):
                    rows_data.append([c.value for c in row])
                else:
                    rows_data.append([row.value])
        else:
            rows_data.append([cells.value])

        headers = [str(val or "") for val in rows_data[0]] if rows_data else []
        return {
            "sheet": ws.title,
            "range": range_str,
            "headers": headers,
            "rows": rows_data[1:] if len(rows_data) > 1 else [],
            "all_values": rows_data,
        }

    def read_sheet(self, session_id: Optional[str] = None, sheet: Optional[str] = None, max_rows: int = 100) -> dict[str, Any]:
        """Read structured contents of an entire sheet up to max_rows."""
        _, wb = self._get_wb(session_id)
        ws = self._get_ws(wb, sheet)

        rows_data = []
        for row in ws.iter_rows(max_row=max_rows, values_only=True):
            if any(cell is not None for cell in row):
                rows_data.append(list(row))

        headers = [str(val or "") for val in rows_data[0]] if rows_data else []
        return {
            "sheet": ws.title,
            "headers": headers,
            "total_rows_read": len(rows_data),
            "rows": rows_data[1:] if len(rows_data) > 1 else [],
        }

    def write_cell(self, session_id: Optional[str], cell: str, value: Any, sheet: Optional[str] = None) -> str:
        """Write a value to a specific cell."""
        sid, wb = self._get_wb(session_id)
        ws = self._get_ws(wb, sheet)
        ws[cell] = value
        self._sessions[sid]["is_dirty"] = True
        return sid

    def write_range(
        self,
        session_id: Optional[str],
        start_cell: str,
        values: list[list[Any]],
        sheet: Optional[str] = None,
    ) -> str:
        """Write a 2D matrix of values starting at start_cell (e.g. 'A1')."""
        sid, wb = self._get_wb(session_id)
        ws = self._get_ws(wb, sheet)

        start_col_char = "".join([c for c in start_cell if c.isalpha()]).upper() or "A"
        start_row_num = int("".join([c for c in start_cell if c.isdigit()]) or "1")

        start_col_idx = openpyxl.utils.column_index_from_string(start_col_char)

        for r_idx, row_values in enumerate(values):
            for c_idx, val in enumerate(row_values):
                cell_obj = ws.cell(row=start_row_num + r_idx, column=start_col_idx + c_idx)
                # Auto convert numeric strings
                if isinstance(val, str) and val.replace(".", "", 1).replace("-", "", 1).isdigit():
                    cell_obj.value = float(val) if "." in val else int(val)
                else:
                    cell_obj.value = val

        self._sessions[sid]["is_dirty"] = True
        return sid

    def add_formula(self, session_id: Optional[str], cell: str, formula: str, sheet: Optional[str] = None) -> str:
        """Write an Excel formula string to a cell (e.g. '=SUM(D2:D100)')."""
        sid, wb = self._get_wb(session_id)
        ws = self._get_ws(wb, sheet)
        formula_str = formula if formula.startswith("=") else f"={formula}"
        ws[cell] = formula_str
        self._sessions[sid]["is_dirty"] = True
        return sid

    def format_range(
        self,
        session_id: Optional[str],
        range_str: str,
        font_name: Optional[str] = None,
        font_size: Optional[float] = None,
        bold: Optional[bool] = None,
        italic: Optional[bool] = None,
        color: Optional[str] = None,
        fill_color: Optional[str] = None,
        align_horizontal: Optional[str] = None,
        number_format: Optional[str] = None,
        border: Optional[bool] = None,
        sheet: Optional[str] = None,
    ) -> str:
        """Apply formatting styles to a range of cells."""
        sid, wb = self._get_wb(session_id)
        ws = self._get_ws(wb, sheet)

        font_kwargs: dict[str, Any] = {}
        if font_name: font_kwargs["name"] = font_name
        if font_size: font_kwargs["size"] = font_size
        if bold is not None: font_kwargs["bold"] = bold
        if italic is not None: font_kwargs["italic"] = italic
        if color: font_kwargs["color"] = color.replace("#", "")

        font_obj = Font(**font_kwargs) if font_kwargs else None
        fill_obj = PatternFill(start_color=fill_color.replace("#", ""), end_color=fill_color.replace("#", ""), fill_type="solid") if fill_color else None
        align_obj = Alignment(horizontal=align_horizontal) if align_horizontal else None

        border_obj = None
        if border:
            thin = Side(border_style="thin", color="D3D3D3")
            border_obj = Border(left=thin, right=thin, top=thin, bottom=thin)

        # Predefined number formats mapping
        num_fmt_map = {
            "currency": "$#,##0.00",
            "percentage": "0.0%",
            "integer": "#,##0",
            "decimal": "#,##0.00",
            "date": "yyyy-mm-dd",
        }
        fmt_str = num_fmt_map.get((number_format or "").lower(), number_format)

        cells = ws[range_str]
        cell_list = []
        if isinstance(cells, tuple):
            for r in cells:
                if isinstance(r, tuple):
                    cell_list.extend(r)
                else:
                    cell_list.append(r)
        else:
            cell_list.append(cells)

        for cell in cell_list:
            if font_obj: cell.font = font_obj
            if fill_obj: cell.fill = fill_obj
            if align_obj: cell.alignment = align_obj
            if border_obj: cell.border = border_obj
            if fmt_str: cell.number_format = fmt_str

        self._sessions[sid]["is_dirty"] = True
        return sid

    def create_table(
        self,
        session_id: Optional[str],
        range_str: str,
        name: str = "DataTable",
        style: str = "TableStyleMedium9",
        sheet: Optional[str] = None,
    ) -> str:
        """Create a native Excel Table over a range."""
        sid, wb = self._get_wb(session_id)
        ws = self._get_ws(wb, sheet)
        tbl_name = name.replace(" ", "_")
        tab = Table(displayName=tbl_name, ref=range_str)
        tab.tableStyleInfo = TableStyleInfo(name=style, showFirstColumn=False, showLastColumn=False, showRowStripes=True, showColumnStripes=True)
        ws.add_table(tab)
        self._sessions[sid]["is_dirty"] = True
        return sid

    def create_chart(
        self,
        session_id: Optional[str],
        chart_type: str,
        data_range: str,
        title: str = "Chart Summary",
        position: str = "E2",
        categories_range: Optional[str] = None,
        sheet: Optional[str] = None,
    ) -> str:
        """Add a BarChart, LineChart, or PieChart from a data range."""
        sid, wb = self._get_wb(session_id)
        ws = self._get_ws(wb, sheet)

        ctype = chart_type.lower()
        if ctype in ("line", "linechart"):
            chart: openpyxl.chart._chart.ChartBase = LineChart()
        elif ctype in ("pie", "piechart"):
            chart = PieChart()
        else:
            chart = BarChart()

        chart.title = title
        chart.style = 10

        min_col, min_row, max_col, max_row = openpyxl.utils.range_boundaries(data_range)
        data_ref = Reference(ws, min_col=min_col, min_row=min_row, max_col=max_col, max_row=max_row)
        chart.add_data(data_ref, titles_from_data=True)

        if categories_range:
            c_min_col, c_min_row, c_max_col, c_max_row = openpyxl.utils.range_boundaries(categories_range)
            cats_ref = Reference(ws, min_col=c_min_col, min_row=c_min_row, max_col=c_max_col, max_row=c_max_row)
            chart.set_categories(cats_ref)

        ws.add_chart(chart, position)
        self._sessions[sid]["is_dirty"] = True
        return sid

    def save_session(self, session_id: Optional[str] = None, output_path: Optional[str] = None) -> tuple[str, str]:
        """Save the active workbook session to disk."""
        sid, wb = self._get_wb(session_id)
        session = self._sessions[sid]

        target_path: Optional[Path] = None
        if output_path:
            target_path = self._resolve_path(output_path)
        elif session.get("path"):
            target_path = session["path"]
        else:
            from backend.core.paths import get_active_desktop
            target_path = get_active_desktop() / f"{sid}.xlsx"

        target_path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(str(target_path))
        session["path"] = target_path
        session["is_dirty"] = False
        logger.info(f"Saved xlsx session '{sid}' to {target_path}")
        return sid, str(target_path)

    def read_spreadsheet(self, path_or_session_id: str) -> SpreadsheetStats:
        """Inspect and return metrics for an XLSX file path or active session_id."""
        wb: Optional[openpyxl.Workbook] = None
        resolved_path = ""
        sid = ""

        if path_or_session_id in self._sessions:
            sid = path_or_session_id
            session = self._sessions[sid]
            wb = session["wb"]
            resolved_path = str(session["path"]) if session.get("path") else ""
        else:
            path_obj = self._resolve_path(path_or_session_id)
            resolved_path = str(path_obj)
            if not path_obj.exists():
                return SpreadsheetStats(path=resolved_path, exists=False, readable=False)
            try:
                wb = openpyxl.load_workbook(filename=resolved_path, data_only=False)
            except Exception as err:
                logger.error(f"Failed to parse xlsx at {resolved_path}: {err}")
                return SpreadsheetStats(path=resolved_path, exists=True, readable=False)

        sheet_names = list(wb.sheetnames)
        active_sheet_title = wb.active.title if wb.active else (sheet_names[0] if sheet_names else "")

        table_count = 0
        chart_count = 0
        formula_count = 0
        total_cells = 0
        dims = {}

        for sheet_name in sheet_names:
            ws = wb[sheet_name]
            table_count += len(ws._tables)
            chart_count += len(ws._charts)
            dims[sheet_name] = ws.dimensions or "A1"

            for row in ws.iter_rows(values_only=False):
                for cell in row:
                    if cell.value is not None:
                        total_cells += 1
                        if isinstance(cell.value, str) and cell.value.startswith("="):
                            formula_count += 1

        return SpreadsheetStats(
            path=resolved_path,
            session_id=sid,
            sheet_names=sheet_names,
            active_sheet=active_sheet_title,
            total_sheets=len(sheet_names),
            tables=table_count,
            charts=chart_count,
            formulas=formula_count,
            total_cells_populated=total_cells,
            exists=True,
            readable=True,
            sheet_dimensions=dims,
        )

    def close_session(self, session_id: str) -> None:
        """Close and remove a session from memory."""
        if session_id in self._sessions:
            del self._sessions[session_id]
            if self._last_active_session_id == session_id:
                self._last_active_session_id = next(iter(self._sessions.keys()), None)


# Global adapter instance for NEXUS execution loop
excel_adapter = ExcelAdapter()
