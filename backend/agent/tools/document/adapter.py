"""Adapter encapsulating python-docx operations and document session management."""
from __future__ import annotations
import os
import uuid
import logging
from pathlib import Path
from typing import Any, Optional

import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

from backend.agent.tools.document.models import DocumentStats

logger = logging.getLogger("nexus.docx_adapter")


class DocxAdapter:
    """Stateful adapter for managing DOCX documents and sessions via python-docx."""

    def __init__(self):
        # Maps session_id -> {"doc": Document, "path": Path | None, "is_dirty": bool}
        self._sessions: dict[str, dict[str, Any]] = {}
        # Keeps track of the most recently active session ID
        self._last_active_session_id: Optional[str] = None

    def _resolve_path(self, raw_path: str) -> Path:
        """Expand user, env vars, handle OneDrive redirection, and relative shortcut prefixes (e.g. Desktop/file.docx)."""
        from backend.core.paths import resolve_system_path
        return resolve_system_path(raw_path, default_filename="document.docx")

    def _resolve_session_id(self, session_id: Optional[str] = None) -> str:
        """Resolve a session_id, falling back to the last active session or default session."""
        if session_id and session_id in self._sessions:
            return session_id
        if self._last_active_session_id and self._last_active_session_id in self._sessions:
            return self._last_active_session_id
        if self._sessions:
            return next(iter(self._sessions.keys()))
        # Auto-create a default session if none exists
        new_id, _ = self.create_session()
        return new_id

    def create_session(self, path: Optional[str] = None, session_id: Optional[str] = None) -> tuple[str, str]:
        """Create a new blank in-memory DOCX document session."""
        sid = session_id or f"doc_{uuid.uuid4().hex[:8]}"
        doc = docx.Document()
        target_path = self._resolve_path(path) if path else None

        self._sessions[sid] = {
            "doc": doc,
            "path": target_path,
            "is_dirty": True,
        }
        self._last_active_session_id = sid
        resolved_path_str = str(target_path) if target_path else ""
        logger.info(f"Created docx session '{sid}' (path: {resolved_path_str})")
        return sid, resolved_path_str

    def open_session(self, path: str, session_id: Optional[str] = None) -> tuple[str, str]:
        """Open an existing .docx file into a session."""
        target_path = self._resolve_path(path)
        if not target_path.exists():
            raise FileNotFoundError(f"Document file not found: {target_path}")

        doc = docx.Document(str(target_path))
        sid = session_id or f"doc_{uuid.uuid4().hex[:8]}"
        self._sessions[sid] = {
            "doc": doc,
            "path": target_path,
            "is_dirty": False,
        }
        self._last_active_session_id = sid
        logger.info(f"Opened docx session '{sid}' from {target_path}")
        return sid, str(target_path)

    def _get_doc(self, session_id: Optional[str] = None) -> tuple[str, docx.Document]:
        sid = self._resolve_session_id(session_id)
        session = self._sessions.get(sid)
        if not session or "doc" not in session:
            raise KeyError(f"Document session '{sid}' not found or already closed.")
        return sid, session["doc"]

    def add_title(self, session_id: Optional[str], text: str) -> str:
        """Add a formatted title paragraph to the document."""
        sid, doc = self._get_doc(session_id)
        p = doc.add_paragraph()
        run = p.add_run(text)
        run.font.size = Pt(24)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x78)  # Deep professional blue
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after = Pt(12)
        self._sessions[sid]["is_dirty"] = True
        return sid

    def add_heading(self, session_id: Optional[str], text: str, level: int = 1) -> str:
        """Add a heading with a specified level (1 to 4)."""
        sid, doc = self._get_doc(session_id)
        clamped_level = max(1, min(4, int(level)))
        doc.add_heading(text, level=clamped_level)
        self._sessions[sid]["is_dirty"] = True
        return sid

    def add_paragraph(self, session_id: Optional[str], text: str, style: str = "Normal") -> str:
        """Add a paragraph of text."""
        sid, doc = self._get_doc(session_id)
        doc.add_paragraph(text, style=style)
        self._sessions[sid]["is_dirty"] = True
        return sid

    def add_bullet(self, session_id: Optional[str], text: str) -> str:
        """Add a bullet point list item."""
        sid, doc = self._get_doc(session_id)
        doc.add_paragraph(text, style="List Bullet")
        self._sessions[sid]["is_dirty"] = True
        return sid

    def add_numbered_item(self, session_id: Optional[str], text: str) -> str:
        """Add a numbered list item."""
        sid, doc = self._get_doc(session_id)
        doc.add_paragraph(text, style="List Number")
        self._sessions[sid]["is_dirty"] = True
        return sid

    def add_table(
        self,
        session_id: Optional[str],
        headers: list[str],
        rows: list[list[str]],
        style: str = "Table Grid",
    ) -> str:
        """Add a structured table with header row and data rows."""
        sid, doc = self._get_doc(session_id)
        num_cols = max(len(headers), max((len(r) for r in rows), default=0))
        if num_cols == 0:
            raise ValueError("Table must have at least 1 column or row.")

        table = doc.add_table(rows=1, cols=num_cols)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        try:
            table.style = style
        except Exception:
            table.style = "Table Grid"

        # Populate header row
        hdr_cells = table.rows[0].cells
        for idx, header_text in enumerate(headers):
            if idx < num_cols:
                cell = hdr_cells[idx]
                cell.text = str(header_text)
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.bold = True

        # Populate data rows
        for row_data in rows:
            row_cells = table.add_row().cells
            for idx, cell_value in enumerate(row_data):
                if idx < num_cols:
                    row_cells[idx].text = str(cell_value)

        doc.add_paragraph()  # Spacing paragraph after table
        self._sessions[sid]["is_dirty"] = True
        return sid

    def add_page_break(self, session_id: Optional[str]) -> str:
        """Add a page break."""
        sid, doc = self._get_doc(session_id)
        doc.add_page_break()
        self._sessions[sid]["is_dirty"] = True
        return sid

    def add_image(self, session_id: Optional[str], image_path: str, width_inches: Optional[float] = None) -> str:
        """Add an image from local filesystem path."""
        sid, doc = self._get_doc(session_id)
        img_p = Path(image_path).resolve()
        if not img_p.exists():
            raise FileNotFoundError(f"Image file not found: {img_p}")

        if width_inches:
            doc.add_picture(str(img_p), width=Inches(width_inches))
        else:
            doc.add_picture(str(img_p), width=Inches(6.0))
        self._sessions[sid]["is_dirty"] = True
        return sid

    def save_session(self, session_id: Optional[str] = None, output_path: Optional[str] = None) -> tuple[str, str]:
        """Save the document session to disk."""
        sid, doc = self._get_doc(session_id)
        session = self._sessions[sid]

        target_path: Optional[Path] = None
        if output_path:
            target_path = self._resolve_path(output_path)
        elif session.get("path"):
            target_path = session["path"]
        else:
            from backend.core.paths import get_active_documents
            target_path = get_active_documents() / f"{sid}.docx"

        target_path.parent.mkdir(parents=True, exist_ok=True)
        doc.save(str(target_path))
        session["path"] = target_path
        session["is_dirty"] = False
        logger.info(f"Saved document session '{sid}' to {target_path}")
        return sid, str(target_path)

    def read_document(self, path_or_session_id: str) -> DocumentStats:
        """Inspect and return metrics for a file path or active session_id."""
        doc: Optional[docx.Document] = None
        resolved_path = ""
        sid = ""

        if not path_or_session_id:
            if self._last_active_session_id and self._last_active_session_id in self._sessions:
                path_or_session_id = self._last_active_session_id
            elif self._sessions:
                path_or_session_id = next(iter(self._sessions.keys()))

        if path_or_session_id in self._sessions:
            sid = path_or_session_id
            session = self._sessions[sid]
            doc = session["doc"]
            resolved_path = str(session["path"]) if session.get("path") else ""
        else:
            path_obj = self._resolve_path(path_or_session_id)
            resolved_path = str(path_obj)
            if not path_obj.exists():
                return DocumentStats(path=resolved_path, exists=False, readable=False)
            try:
                doc = docx.Document(resolved_path)
            except Exception as err:
                logger.error(f"Failed to parse docx at {resolved_path}: {err}")
                return DocumentStats(path=resolved_path, exists=True, readable=False)

        para_count = len(doc.paragraphs)
        headings_count = 0
        heading_titles = []
        char_count = 0

        for p in doc.paragraphs:
            text = p.text.strip()
            char_count += len(text)
            if p.style and p.style.name.startswith("Heading"):
                headings_count += 1
                heading_titles.append(text)

        table_count = len(doc.tables)
        table_dims = []
        for t in doc.tables:
            table_dims.append({"rows": len(t.rows), "cols": len(t.columns)})
            for row in t.rows:
                for cell in row.cells:
                    char_count += len(cell.text.strip())

        section_count = len(doc.sections)

        return DocumentStats(
            path=resolved_path,
            session_id=sid,
            paragraphs=para_count,
            headings=headings_count,
            tables=table_count,
            characters=char_count,
            sections=section_count,
            exists=True,
            readable=True,
            heading_titles=heading_titles,
            table_dimensions=table_dims,
        )

    def close_session(self, session_id: str) -> None:
        """Close and remove a session from memory."""
        if session_id in self._sessions:
            del self._sessions[session_id]
            if self._last_active_session_id == session_id:
                self._last_active_session_id = next(iter(self._sessions.keys()), None)


# Global adapter instance for NEXUS execution loop
docx_adapter = DocxAdapter()
